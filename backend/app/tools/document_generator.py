import os
from pathlib import Path
from typing import Any
from datetime import datetime, timezone
from app.config import settings
from app.security.path_validator import validate_and_resolve_path

class DocumentGenerationError(Exception):
    pass

def generate_pptx_file(title: str, subtitle: str, slides: list[dict[str, Any]], filename: str) -> dict[str, Any]:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    if not filename.endswith(".pptx"):
        filename += ".pptx"

    target_path = validate_and_resolve_path(str(settings.OUTPUT_DIR / filename), allow_output=True)
    prs = Presentation()

    DARK_BLUE = RGBColor(13, 17, 39)
    ACCENT_BLUE = RGBColor(59, 130, 246)
    TEXT_MUTED = RGBColor(148, 163, 184)

    blank_slide_layout = prs.slide_layouts[6]
    title_slide = prs.slides.add_slide(blank_slide_layout)
    
    tx_box = title_slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
    tf = tx_box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(20)
        p2.font.color.rgb = TEXT_MUTED

    for slide_data in slides:
        slide = prs.slides.add_slide(blank_slide_layout)
        
        h_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(8.4), Inches(1))
        h_tf = h_box.text_frame
        h_p = h_tf.paragraphs[0]
        h_p.text = slide_data.get("title", "Topic")
        h_p.font.size = Pt(28)
        h_p.font.bold = True
        h_p.font.color.rgb = ACCENT_BLUE

        bullets = slide_data.get("bullets", [])
        if bullets:
            b_box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(4.5))
            b_tf = b_box.text_frame
            b_tf.word_wrap = True
            for i, bullet in enumerate(bullets):
                bp = b_tf.paragraphs[0] if i == 0 else b_tf.add_paragraph()
                bp.text = f"•  {bullet}"
                bp.font.size = Pt(18)
                bp.font.color.rgb = DARK_BLUE
                bp.space_after = Pt(12)

    prs.save(str(target_path))

    if not target_path.exists() or target_path.stat().st_size == 0:
        raise DocumentGenerationError(f"File verification failed: {target_path} not found or empty.")

    return {
        "status": "success",
        "file_type": "pptx",
        "filename": target_path.name,
        "path": str(target_path),
        "size_bytes": target_path.stat().st_size,
        "slides_count": len(slides) + 1,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }


def generate_docx_file(title: str, sections: list[dict[str, Any]], filename: str) -> dict[str, Any]:
    from docx import Document
    from docx.shared import Pt, RGBColor

    if not filename.endswith(".docx"):
        filename += ".docx"

    target_path = validate_and_resolve_path(str(settings.OUTPUT_DIR / filename), allow_output=True)
    doc = Document()

    title_p = doc.add_paragraph()
    run = title_p.add_run(title)
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = RGBColor(13, 17, 39)
    title_p.paragraph_format.space_after = Pt(18)

    for sec in sections:
        sec_title = sec.get("heading")
        if sec_title:
            h = doc.add_heading(sec_title, level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)

        content = sec.get("content")
        if content:
            p = doc.add_paragraph(content)
            p.paragraph_format.space_after = Pt(8)

        bullets = sec.get("bullets", [])
        for b in bullets:
            bp = doc.add_paragraph(b, style="List Bullet")
            bp.paragraph_format.space_after = Pt(4)

    doc.save(str(target_path))

    if not target_path.exists() or target_path.stat().st_size == 0:
        raise DocumentGenerationError(f"Word document verification failed: {target_path}")

    return {
        "status": "success",
        "file_type": "docx",
        "filename": target_path.name,
        "path": str(target_path),
        "size_bytes": target_path.stat().st_size,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }


def generate_pdf_file(title: str, sections: list[dict[str, Any]], filename: str) -> dict[str, Any]:
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors

    if not filename.endswith(".pdf"):
        filename += ".pdf"

    target_path = validate_and_resolve_path(str(settings.OUTPUT_DIR / filename), allow_output=True)

    doc = SimpleDocTemplate(
        str(target_path),
        pagesize=letter,
        rightMargin=54,
        leftMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=24,
        leading=28,
        textColor=colors.HexColor("#0f172a"),
        spaceAfter=12
    )
    heading_style = ParagraphStyle(
        'DocHeading',
        parent=styles['Heading2'],
        fontSize=15,
        leading=18,
        textColor=colors.HexColor("#2563eb"),
        spaceBefore=14,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#334155"),
        spaceAfter=6
    )
    bullet_style = ParagraphStyle(
        'DocBullet',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#1e293b"),
        leftIndent=15,
        spaceAfter=4
    )

    story = [
        Paragraph(title, title_style),
        HRFlowable(width="100%", thickness=1, color=colors.HexColor("#cbd5e1"), spaceAfter=14)
    ]

    for sec in sections:
        h = sec.get("heading")
        if h:
            story.append(Paragraph(h, heading_style))

        c = sec.get("content")
        if c:
            story.append(Paragraph(c, body_style))

        for b in sec.get("bullets", []):
            story.append(Paragraph(f"&bull; {b}", bullet_style))

        story.append(Spacer(1, 8))

    doc.build(story)

    if not target_path.exists() or target_path.stat().st_size == 0:
        raise DocumentGenerationError(f"PDF generation verification failed: {target_path}")

    return {
        "status": "success",
        "file_type": "pdf",
        "filename": target_path.name,
        "path": str(target_path),
        "size_bytes": target_path.stat().st_size,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
